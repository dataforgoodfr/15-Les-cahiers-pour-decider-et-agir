"""Téléchargement et lecture des référentiels de l'INSEE.

Seul module qui accède au réseau. Chaque fichier est gardé en cache local
(`data/sources/` par défaut) : une seconde exécution ne retélécharge rien.
"""

import csv
import io
import urllib.request
import zipfile
from pathlib import Path

import openpyxl

# Code officiel géographique au 1er janvier 2019 : https://www.insee.fr/fr/information/3720946
COG_COMMUNES = (
    "https://www.insee.fr/fr/statistiques/fichier/3720946/communes-01012019-csv.zip"
)
COG_DEPARTEMENTS = (
    "https://www.insee.fr/fr/statistiques/fichier/3720946/departement2019-csv.zip"
)
COG_REGIONS = "https://www.insee.fr/fr/statistiques/fichier/3720946/region2019-csv.zip"
# Populations légales millésimées 2017, dans les limites au 1er janvier 2019 :
# https://www.insee.fr/fr/statistiques/4265429
POPULATIONS = "https://www.insee.fr/fr/statistiques/fichier/4265429/ensemble.zip"
# Mouvements de communes depuis 1943 : https://www.insee.fr/fr/information/8740222
MOUVEMENTS = (
    "https://www.insee.fr/fr/statistiques/fichier/8740222/v_mvt_commune_2026.csv"
)
# Grille communale de densité à 4 niveaux, géographie au 1er janvier 2021 :
# https://www.insee.fr/fr/information/2114627
GRILLE_DENSITE = (
    "https://www.insee.fr/fr/statistiques/fichier/2114627/grille_densite_2021.zip"
)
# Recensement 2017, structure de la population par commune (âges, CSP),
# géographie au 1er janvier 2019 : https://www.insee.fr/fr/statistiques/4515565
STRUCTURE_POPULATION = (
    "https://www.insee.fr/fr/statistiques/fichier/4515565/"
    "base-ccc-evol-struct-pop-2017.zip"
)
# Filosofi 2017, revenu disponible par commune, géographie au 1er janvier 2018 :
# https://www.insee.fr/fr/statistiques/4291712
REVENUS = (
    "https://www.insee.fr/fr/statistiques/fichier/4291712/"
    "indic-struct-distrib-revenu-2017-COMMUNES.zip"
)


def telecharger(url: str, cache: Path) -> bytes:
    """Récupère une URL, ou sa copie locale si elle existe déjà."""
    fichier = cache / url.rsplit("/", 1)[-1]
    if fichier.exists():
        return fichier.read_bytes()
    requete = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(requete, timeout=180) as reponse:
        contenu = reponse.read()
    cache.mkdir(parents=True, exist_ok=True)
    fichier.write_bytes(contenu)
    return contenu


def lignes_csv(texte: str, delimiteur: str = ",") -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(texte.lstrip("﻿")), delimiter=delimiteur))


def membre_zip(contenu: bytes, nom: str) -> str:
    with zipfile.ZipFile(io.BytesIO(contenu)) as archive:
        return archive.read(nom).decode("utf-8")


def premier_csv_zip(contenu: bytes) -> list[dict[str, str]]:
    """Lit l'unique CSV d'une archive du COG."""
    with zipfile.ZipFile(io.BytesIO(contenu)) as archive:
        nom = next(n for n in archive.namelist() if n.endswith(".csv"))
        return lignes_csv(archive.read(nom).decode("utf-8"))


def grille_densite(contenu: bytes) -> dict[str, int]:
    """Degré de densité (1 à 4) par code commune, depuis le fichier détaillé."""
    with zipfile.ZipFile(io.BytesIO(contenu)) as archive:
        nom = next(n for n in archive.namelist() if "detaille" in n)
        classeur = openpyxl.load_workbook(io.BytesIO(archive.read(nom)), read_only=True)
    feuille = classeur.worksheets[0]
    grille = {}
    for ligne in feuille.iter_rows(min_row=2, values_only=True):
        if ligne[0] and ligne[2]:
            grille[str(ligne[0]).strip()] = int(ligne[2])
    return grille


def colonnes_csv_zip(
    contenu: bytes, nom: str, colonnes: list[str], delimiteur: str = ";"
) -> dict[str, dict[str, str]]:
    """Quelques colonnes d'un gros CSV d'archive, indexées par CODGEO."""
    with zipfile.ZipFile(io.BytesIO(contenu)) as archive, archive.open(nom) as brut:
        lecteur = csv.DictReader(
            io.TextIOWrapper(brut, encoding="utf-8-sig"), delimiter=delimiteur
        )
        return {ligne["CODGEO"]: {c: ligne[c] for c in colonnes} for ligne in lecteur}


def revenus_medians(contenu: bytes) -> dict[str, int]:
    """Revenu disponible médian par unité de consommation, par code commune.

    Absent (secret statistique) pour les communes de moins de 50 ménages.
    """
    with zipfile.ZipFile(io.BytesIO(contenu)) as archive:
        classeur = openpyxl.load_workbook(
            io.BytesIO(archive.read("FILO2017_DISP_COM.xlsx")), read_only=True
        )
    medianes = {}
    for ligne in classeur["ENSEMBLE"].iter_rows(min_row=7, values_only=True):
        code, mediane = ligne[0], ligne[6]
        if code and isinstance(mediane, (int, float)):
            medianes[str(code)] = int(mediane)
    return medianes
